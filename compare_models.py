"""
Compare RF, XGBoost, LightGBM, CatBoost and a soft-voting ensemble on the
pipeline's 11 features, and write report-ready figures, tables and a report.

    python compare_models.py                  # full run (writes models/ and results/)
    python compare_models.py --quick --out <dir> --models-dir <dir>   # smoke test

Protocol (same features, data, splits and seeds for every model):
  * Data: dsp.snippet / dsp.capture synthetic generators, features from dsp.features
    (the pipeline's own code). SNR is in-band (dsp.SNR_DEFINITION).
  * Train: in-band SNR ~ U(0, 30) dB. Test: fixed SNR steps -10..+30 dB (2 dB) plus
    -5 and +5 dB, N per class per step. Separate RNG streams for train/test/end-to-end.
  * Seeds 0, 1, 2: each draws its own train/test/end-to-end sets and model seed.
  * Tuning: 30-trial random search, 5-fold stratified CV, macro-F1, on the seed-0
    training set only; the best params are reused for seeds 1 and 2. The ensemble
    reuses its members' params; its members are calibrated (CalibratedClassifierCV)
    before soft voting, and an uncalibrated ensemble is evaluated for comparison.
  * Two evaluations: classifier-only (features from the true signal slice) and
    end-to-end (signal placed in a 0.5 s capture, run through detect_regions; success
    = detected AND correctly labelled).
  * The deployed RF (rf_classifier_model.joblib, trained by train_rf_classifier.py)
    is evaluated untouched as a reference.

Outputs: models/<name>.joblib + .json; results/models/ (CSVs, raw plots);
results/report/figures/ (fig*.png + .pdf), results/report/tables/ (csv, md, tex),
results/report/model_comparison_report.md + .docx.
"""

import argparse
import csv
import glob
import json
import math
import os
import platform
import re
import sys
import time
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import joblib
import numpy as np
from joblib import Parallel, delayed
from matplotlib.figure import Figure
from sklearn.metrics import (balanced_accuracy_score, f1_score, log_loss,
                             precision_recall_fscore_support, roc_curve)
from sklearn.model_selection import RandomizedSearchCV, StratifiedKFold, cross_val_score
from sklearn.preprocessing import LabelEncoder

import dsp
from models import BUILDERS, PARAM_SPACES, LabeledModel, build_ensemble, save_model
from plot_style import (CLASS_ORDER, DOUBLE_COL, MODEL_COLORS, MODEL_LABELS, MODEL_ORDER,
                        SINGLE_COL, apply_style, save_figure)

FULL = {"seeds": [0, 1, 2], "train_per_class": 200, "test_per_step": 10, "e2e_per_step": 4,
        "snrs": sorted(set(range(-10, 31, 2)) | {-5, 5}), "trials": 30, "folds": 5}
QUICK = {"seeds": [0, 1], "train_per_class": 20, "test_per_step": 2, "e2e_per_step": 1,
         "snrs": [-10, 0, 5, 10, 20], "trials": 3, "folds": 3}
TABLE_SNRS = [-5, 0, 5, 10, 20]
DATA_VERSION = 2   # bump when dsp's generators/features change, so cached datasets are rebuilt
POSITIVE = "LORA"                 # the drone-link class for Pd at Pfa = 1 %
BLOCK_BUDGET_MS = 500
CENTER_FREQ = 866e6
SAVED_NAME = {"rf": "rf_tuned", "xgb": "xgb", "lgbm": "lgbm", "cat": "cat",
              "ensemble": "ensemble", "ensemble_uncal": "ensemble_uncal"}
SHORT_CLASS = {"FM_BROADCAST": "FM", "DRONE_TELEMETRY": "TELEM", "DRONE_VIDEO": "VIDEO",
               "INTERFERENCE": "CW", "NOISE": "NOISE", "LORA": "LORA"}


# ---------------------------------------------------------------- data

def make_train(seed, cfg):
    jobs = [(seed, ci, i) for ci in range(len(dsp.CLASSES)) for i in range(cfg["train_per_class"])]
    X = Parallel(n_jobs=-1, batch_size=4)(delayed(dsp.train_row)(*j) for j in jobs)
    return np.array(X), np.array([dsp.CLASSES[ci] for _, ci, _ in jobs])


def make_test(seed, cfg):
    jobs = [(seed, s, ci, i) for s in cfg["snrs"] for ci in range(len(dsp.CLASSES))
            for i in range(cfg["test_per_step"])]
    X = Parallel(n_jobs=-1, batch_size=4)(delayed(dsp.test_row)(*j) for j in jobs)
    return (np.array(X), np.array([dsp.CLASSES[j[2]] for j in jobs]),
            np.array([j[1] for j in jobs], dtype=float))


def make_e2e(seed, cfg):
    jobs = [(seed, s, ci, i) for s in cfg["snrs"] for ci in range(len(dsp.CLASSES))
            for i in range(cfg["e2e_per_step"])]
    rows = Parallel(n_jobs=-1, batch_size=2)(delayed(dsp.e2e_row)(*j) for j in jobs)
    for row, (_, s, ci, _) in zip(rows, jobs):
        row["cls"], row["snr"] = dsp.CLASSES[ci], s
    return rows


# ---------------------------------------------------------------- metrics

def aligned_proba(model, X, classes):
    """predict_proba with columns in `classes` order, whatever the model's own order."""
    proba = model.predict_proba(X)
    order = [list(model.classes_).index(c) for c in classes]
    return proba[:, order]


def ece_score(y, proba, classes, bins=15):
    conf = proba.max(axis=1)
    correct = np.array(classes)[proba.argmax(axis=1)] == y
    edges = np.linspace(0, 1, bins + 1)
    ece = 0.0
    for lo, hi in zip(edges[:-1], edges[1:]):
        sel = (conf > lo) & (conf <= hi)
        if sel.any():
            ece += sel.mean() * abs(correct[sel].mean() - conf[sel].mean())
    return ece


def pd_at_pfa(y, proba, classes, pfa=0.01):
    s = proba[:, classes.index(POSITIVE)]
    neg, pos = s[y != POSITIVE], s[y == POSITIVE]
    thr = np.quantile(neg, 1 - pfa)
    return float(np.mean(pos > thr))


def snr_for_f1(snrs, f1s, target=0.9):
    """Lowest SNR where the macro-F1 curve reaches target (linear interpolation), or nan."""
    for k in range(len(snrs)):
        if f1s[k] >= target:
            if k == 0:
                return float(snrs[0])
            s0, s1, f0, f1 = snrs[k - 1], snrs[k], f1s[k - 1], f1s[k]
            return float(s0 + (target - f0) * (s1 - s0) / (f1 - f0))
    return float("nan")


def evaluate(model, X, y, snr, classes, cfg):
    proba = aligned_proba(model, X, classes)
    pred = np.array(classes)[proba.argmax(axis=1)]
    onehot = (np.array(classes)[None, :] == y[:, None]).astype(float)
    p, r, f, _ = precision_recall_fscore_support(y, pred, labels=classes, zero_division=0)
    per_snr = [f1_score(y[snr == s], pred[snr == s], labels=classes, average="macro", zero_division=0)
               for s in cfg["snrs"]]
    band = (snr >= 0) & (snr <= 10)
    return {
        "macro_f1": f1_score(y, pred, labels=classes, average="macro", zero_division=0),
        "balanced_accuracy": balanced_accuracy_score(y, pred),
        "pd_at_pfa1": pd_at_pfa(y, proba, classes),
        "log_loss": log_loss(y, np.clip(proba, 1e-15, 1), labels=classes),
        "brier": float(np.mean(np.sum((proba - onehot) ** 2, axis=1))),
        "ece": ece_score(y, proba, classes),
        "f1_0_10": f1_score(y[band], pred[band], labels=classes, average="macro", zero_division=0),
        "snr_f1_90": snr_for_f1(cfg["snrs"], per_snr),
        "f1_vs_snr": per_snr,
        "precision": dict(zip(classes, p)), "recall": dict(zip(classes, r)), "f1": dict(zip(classes, f)),
        "proba": proba, "pred": pred,
    }


def evaluate_e2e(model, rows, cfg):
    """End-to-end success rate (detected AND correctly labelled) per SNR, over signal classes."""
    feats = [f for row in rows for f in ([row["features"]] if row["cls"] != "NOISE" and row["features"]
                                         else row["features"] if row["cls"] == "NOISE" else [])]
    labels = iter(model.predict(np.array(feats)) if feats else [])
    for row in rows:
        if row["cls"] == "NOISE":
            preds = [next(labels) for _ in row["features"]]
            row["ok"] = all(p == "NOISE" for p in preds)
        else:
            row["ok"] = bool(row["features"]) and next(labels) == row["cls"]
    curve, pd_curve = [], []
    for s in cfg["snrs"]:
        sig = [r for r in rows if r["snr"] == s and r["cls"] != "NOISE"]
        curve.append(np.mean([r["ok"] for r in sig]))
        pd_curve.append(np.mean([r["detected"] for r in sig]))
    noise = [r for r in rows if r["cls"] == "NOISE"]
    band = [r for r in rows if r["cls"] != "NOISE" and 0 <= r["snr"] <= 10]
    high = {c: float(np.mean([r["ok"] for r in rows if r["cls"] == c and r["snr"] >= 20]))
            for c in CLASS_ORDER if c != "NOISE"}
    return {"e2e_vs_snr": curve, "pd_vs_snr": pd_curve, "e2e_high_per_class": high,
            "e2e_0_10": float(np.mean([r["ok"] for r in band])),
            "noise_false_alarm": float(np.mean([not r["ok"] for r in noise]))}


def permutation_drop(model, X, y, repeats=5, seed=0):
    """Mean macro-F1 drop when each feature column is shuffled (permutation importance)."""
    rng = np.random.default_rng(seed)
    base = f1_score(y, model.predict(X), average="macro")
    drops = []
    for j in range(X.shape[1]):
        scores = []
        for _ in range(repeats):
            Xp = X.copy()
            Xp[:, j] = rng.permutation(Xp[:, j])
            scores.append(f1_score(y, model.predict(Xp), average="macro"))
        drops.append(base - float(np.mean(scores)))
    return np.array(drops)


def latency_ms(model, X):
    """Median predict_proba time per detection for batches of 1 and 10."""
    out = {}
    for batch, reps in ((1, 200), (10, 100)):
        xb = X[:batch]
        model.predict_proba(xb)
        times = []
        for _ in range(reps):
            t0 = time.perf_counter()
            model.predict_proba(xb)
            times.append((time.perf_counter() - t0) * 1000 / batch)
        out[f"latency_b{batch}"] = float(np.median(times))
    return out


# ---------------------------------------------------------------- formatting / tables

def sig3(x):
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return "n/a"
    if x == 0:
        return "0"
    digits = max(0, 3 - int(math.floor(math.log10(abs(x)))) - 1)
    return f"{round(x, digits):.{digits}f}"


def pm(mean, std):
    if mean is None or (isinstance(mean, float) and math.isnan(mean)):
        return "n/a"
    m = sig3(mean)
    decimals = len(m.split(".")[1]) if "." in m else 0
    return f"{m} ± {std:.{decimals}f}" if std is not None else m


def write_table(name, title, columns, rows, out_dir, group_key=None):
    """
    Write <name>.csv, .md and .tex (booktabs). columns: list of (key, header, best) with
    best in {"max", "min", None}; row values are str, number or (mean, std). Best value per
    column (within each group_key group, if given) is bold in md/tex.
    """
    os.makedirs(out_dir, exist_ok=True)

    def mean_of(v):
        return v[0] if isinstance(v, tuple) else v

    groups = {}
    for i, row in enumerate(rows):
        groups.setdefault(row.get(group_key) if group_key else None, []).append(i)
    bold = set()
    for key, _, best in columns:
        if best is None:
            continue
        for idx in groups.values():
            vals = [(mean_of(rows[i][key]), i) for i in idx
                    if isinstance(mean_of(rows[i][key]), (int, float)) and not math.isnan(mean_of(rows[i][key]))]
            if vals:
                target = max(vals)[0] if best == "max" else min(vals)[0]
                bold.update((i, key) for v, i in vals if v == target)

    def cell(i, key):
        v = rows[i][key]
        if isinstance(v, tuple):
            return pm(*v)
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return sig3(float(v))
        return str(v)

    csv_path = os.path.join(out_dir, name + ".csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        header = []
        for key, head, _ in columns:
            if any(isinstance(r[key], tuple) for r in rows):
                header += [head + " mean", head + " std"]
            else:
                header.append(head)
        w.writerow(header)
        for r in rows:
            line = []
            for key, _, _ in columns:
                if any(isinstance(x[key], tuple) for x in rows):
                    v = r[key] if isinstance(r[key], tuple) else (r[key], "")
                    line += [v[0], v[1]]
                else:
                    line.append(r[key])
            w.writerow(line)

    md = [f"**{title}**", "", "| " + " | ".join(h for _, h, _ in columns) + " |",
          "|" + "|".join("---" for _ in columns) + "|"]
    for i in range(len(rows)):
        md.append("| " + " | ".join(f"**{cell(i, k)}**" if (i, k) in bold else cell(i, k)
                                    for k, _, _ in columns) + " |")
    md_path = os.path.join(out_dir, name + ".md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md) + "\n")

    def tex(s):
        s = s.replace("\\", r"\textbackslash{}")
        for a, b in (("&", r"\&"), ("%", r"\%"), ("_", r"\_"), ("#", r"\#"), ("±", r"$\pm$"),
                     ("×", r"$\times$"), ("≥", r"$\geq$"), ("−", "-")):
            s = s.replace(a, b)
        return s

    tex_lines = [r"\begin{table}[htbp]", r"\centering", r"\small",
                 rf"\caption{{{tex(title)}}}", rf"\label{{tab:{name}}}",
                 r"\begin{tabular}{" + "l" * 1 + "r" * (len(columns) - 1) + "}", r"\toprule",
                 " & ".join(tex(h) for _, h, _ in columns) + r" \\", r"\midrule"]
    for i in range(len(rows)):
        tex_lines.append(" & ".join(rf"\textbf{{{tex(cell(i, k))}}}" if (i, k) in bold else tex(cell(i, k))
                                    for k, _, _ in columns) + r" \\")
    tex_lines += [r"\bottomrule", r"\end{tabular}", r"\end{table}"]
    tex_path = os.path.join(out_dir, name + ".tex")
    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("\n".join(tex_lines) + "\n")
    return {"csv": csv_path, "md": md_path, "tex": tex_path, "title": title,
            "columns": columns, "rows": rows, "bold": bold, "cell": cell}


# ---------------------------------------------------------------- figures

def legend_outside(ax, fig, ncol=3):
    handles, labels = ax.get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=ncol, frameon=False,
               bbox_to_anchor=(0.5, -0.02))
    fig.subplots_adjust(bottom=0.3)


def fig_curve(snrs, curves, ylabel, title, path, hline=None):
    """curves: {model: array (seeds x snrs)}; mean line with a min..max band over seeds."""
    fig = Figure(figsize=(DOUBLE_COL, 3.6))
    ax = fig.add_subplot()
    for m in MODEL_ORDER:
        if m not in curves:
            continue
        c = np.array(curves[m])
        ax.plot(snrs, c.mean(axis=0), color=MODEL_COLORS[m], label=MODEL_LABELS[m],
                linestyle="--" if m == "rf_deployed" else "-", marker="o", markersize=3)
        ax.fill_between(snrs, c.min(axis=0), c.max(axis=0), color=MODEL_COLORS[m], alpha=0.15, lw=0)
    if hline is not None:
        ax.axhline(hline, color="black", lw=0.8, ls=":")
    ax.set_xlabel("In-band SNR (dB)")
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.set_ylim(0, 1.02)
    legend_outside(ax, fig)
    return save_figure(fig, path)


def fig_confusions(cms, classes, path):
    names = [m for m in MODEL_ORDER if m in cms]
    ncol = 4
    nrow = math.ceil(len(names) / ncol)
    fig = Figure(figsize=(DOUBLE_COL, 2.4 * nrow + 0.8))
    axes = fig.subplots(nrow, ncol, squeeze=False)
    fig.subplots_adjust(hspace=0.9)
    labels = [SHORT_CLASS[c] for c in classes]
    for k, ax in enumerate(axes.flat):
        if k >= len(names):
            ax.axis("off")
            continue
        m = names[k]
        im = ax.imshow(cms[m], vmin=0, vmax=1, cmap="Blues")
        ax.set_title(MODEL_LABELS[m].replace(" (", "\n("), fontsize=10)
        ax.set_xticks(range(len(classes)), labels, rotation=90, fontsize=10)
        ax.set_yticks(range(len(classes)), labels if k % ncol == 0 else [""] * len(classes), fontsize=10)
        ax.grid(False)
        for i in range(len(classes)):
            for j in range(len(classes)):
                if cms[m][i, j] >= 0.1:
                    ax.text(j, i, f"{cms[m][i, j]:.1f}".lstrip("0"), ha="center", va="center", fontsize=7,
                            color="white" if cms[m][i, j] > 0.6 else "black")
    fig.supxlabel("Predicted class", fontsize=10)
    fig.supylabel("True class", fontsize=10)
    fig.colorbar(im, ax=axes.ravel().tolist(), shrink=0.6, label="Fraction of true class")
    return save_figure(fig, path)


def fig_roc(rocs, path):
    fig = Figure(figsize=(SINGLE_COL + 1.2, 3.4))
    ax = fig.add_subplot()
    for m in MODEL_ORDER:
        if m in rocs:
            fpr, tpr = rocs[m]
            ax.plot(np.maximum(fpr, 1e-4), tpr, color=MODEL_COLORS[m], label=MODEL_LABELS[m],
                    linestyle="--" if m == "rf_deployed" else "-")
    ax.axvline(0.01, color="black", lw=0.8, ls=":")
    ax.text(0.011, 0.05, "Pfa = 1 %", fontsize=10)
    ax.set_xscale("log")
    ax.set_xlim(1e-4, 1)
    ax.set_ylim(0, 1.02)
    ax.set_xlabel("False-alarm probability Pfa (LORA vs rest)")
    ax.set_ylabel("Detection probability Pd")
    legend_outside(ax, fig, ncol=2)
    fig.subplots_adjust(bottom=0.42)
    return save_figure(fig, path)


def fig_reliability(rel, path):
    names = [m for m in MODEL_ORDER if m in rel]
    ncol = 4
    nrow = math.ceil(len(names) / ncol)
    fig = Figure(figsize=(DOUBLE_COL, 2.6 * nrow + 0.5))
    axes = fig.subplots(nrow, ncol, squeeze=False, sharex=True, sharey=True)
    for k, ax in enumerate(axes.flat):
        if k >= len(names):
            ax.axis("off")
            continue
        m = names[k]
        conf, acc, ece = rel[m]
        ax.plot([0, 1], [0, 1], color="black", lw=0.8, ls=":")
        ax.plot(conf, acc, marker="o", markersize=3, color=MODEL_COLORS[m])
        ax.set_title(f"{MODEL_LABELS[m].replace(' (', chr(10) + '(')}\nECE {ece:.3f}", fontsize=10)
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
    fig.supxlabel("Mean predicted confidence (top label)", fontsize=10)
    fig.supylabel("Observed accuracy", fontsize=10)
    fig.tight_layout()
    return save_figure(fig, path)


def fig_latency(points, path):
    fig = Figure(figsize=(SINGLE_COL + 1.2, 3.2))
    ax = fig.add_subplot()
    for m in MODEL_ORDER:
        if m in points:
            x, y = points[m]
            ax.scatter([x], [y], color=MODEL_COLORS[m], s=40, label=MODEL_LABELS[m], zorder=3)
    ax.set_xscale("log")
    ax.set_xlabel("Inference latency, batch of 1 (ms per detection)")
    ax.set_ylabel("Macro-F1, 0–10 dB SNR")
    legend_outside(ax, fig, ncol=2)
    fig.subplots_adjust(bottom=0.45)
    return save_figure(fig, path)


def fig_heatmap(matrix, row_names, col_names, cbar_label, path, fmt="{:.2f}", vmin=None, vmax=None,
                cmap="viridis", height=None):
    fig = Figure(figsize=(DOUBLE_COL, height or 0.4 * len(row_names) + 1.6))
    ax = fig.add_subplot()
    im = ax.imshow(matrix, aspect="auto", cmap=cmap, vmin=vmin, vmax=vmax)
    ax.set_xticks(range(len(col_names)), col_names, rotation=30, ha="right")
    ax.set_yticks(range(len(row_names)), row_names)
    ax.grid(False)
    lo, hi = np.nanmin(matrix), np.nanmax(matrix)
    for i in range(len(row_names)):
        for j in range(len(col_names)):
            v = matrix[i][j]
            ax.text(j, i, fmt.format(v), ha="center", va="center", fontsize=8,
                    color="white" if (v - lo) < 0.5 * (hi - lo + 1e-12) else "black")
    fig.colorbar(im, ax=ax, label=cbar_label)
    fig.tight_layout()
    return save_figure(fig, path)


def fig_bars(names, synth, real, path):
    fig = Figure(figsize=(DOUBLE_COL, 3.0))
    ax = fig.add_subplot()
    x = np.arange(len(names))
    ax.bar(x - 0.2, synth, 0.4, label="Synthetic test", color="#999999")
    ax.bar(x + 0.2, real, 0.4, label="Real recordings", color="#0072B2")
    ax.set_xticks(x, [MODEL_LABELS[m] for m in names], rotation=20, ha="right")
    ax.set_ylabel("Balanced accuracy, LORA vs rest")
    ax.set_ylim(0, 1.02)
    legend_outside(ax, fig, ncol=2)
    return save_figure(fig, path)


# ---------------------------------------------------------------- real data

def session_of(path, root):
    rel = os.path.relpath(path, root).split(os.sep)
    if len(rel) > 2:
        return rel[1]
    m = re.search(r"(\d{2}-\d{2}-\d{4})", os.path.basename(path))
    return m.group(1) if m else "session-1"


def real_eval(models, root):
    """Clip-level LORA-vs-rest decisions on rf_dataset/lora_hop (positive) and rf_dataset/noise."""
    from rf_scanner import load_recording
    clips = [(p, folder == "lora_hop") for folder in ("lora_hop", "noise")
             for p in sorted(glob.glob(os.path.join(root, folder, "**", "*.wav"), recursive=True))]
    if not clips:
        return None
    rows = []
    for path, label in clips:
        iq, sr = load_recording(path)
        m = re.search(r"_(\d{6,})Hz", os.path.basename(path))
        fc = float(m.group(1)) if m else CENTER_FREQ
        regions = dsp.detect_regions(iq, sr, fc)
        feats = [dsp.region_features(iq, sr, fc, t, regions)[4] for t in regions["targets"]]
        row = {"clip": os.path.basename(path), "session": session_of(path, root), "label": label}
        for name, model in models.items():
            row[name] = bool(feats) and POSITIVE in set(model.predict(np.array(feats)))
        rows.append(row)
    results = {}
    for name in models:
        y = np.array([r["label"] for r in rows])
        p = np.array([r[name] for r in rows])
        sessions = {}
        for s in sorted({r["session"] for r in rows}):
            idx = np.array([r["session"] == s for r in rows])
            sessions[s] = float(balanced_accuracy_score(y[idx], p[idx])) if len(set(y[idx])) > 1 \
                else float(np.mean(y[idx] == p[idx]))
        results[name] = {"balanced_accuracy": float(balanced_accuracy_score(y, p)) if len(set(y)) > 1
                         else float(np.mean(y == p)), "sessions": sessions}
    return {"clips": rows, "results": results}


# ---------------------------------------------------------------- report

def environment():
    import catboost, lightgbm, scipy, sklearn, xgboost, matplotlib as mpl
    return {
        "OS": platform.platform(), "CPU": platform.processor() or platform.machine(),
        "Logical CPUs": os.cpu_count(), "Python": platform.python_version(),
        "numpy": np.__version__, "scipy": scipy.__version__, "scikit-learn": sklearn.__version__,
        "xgboost": xgboost.__version__, "lightgbm": lightgbm.__version__,
        "catboost": catboost.__version__, "matplotlib": mpl.__version__,
    }


def render_report(blocks, md_path, docx_path):
    """blocks: ("h1"|"h2"|"p"|"bullets"|"fig"|"table", ...) rendered to Markdown and .docx."""
    md = []
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            md += [f"# {b[1]}", ""]
        elif kind == "h2":
            md += [f"## {b[1]}", ""]
        elif kind == "p":
            md += [b[1], ""]
        elif kind == "bullets":
            md += [f"- {x}" for x in b[1]] + [""]
        elif kind == "fig":
            rel = os.path.relpath(b[1], os.path.dirname(md_path)).replace(os.sep, "/")
            md += [f"![{b[2]}]({rel})", "", f"*{b[2]}*", ""]
        elif kind == "table":
            t = b[1]
            md += [open(t["md"], encoding="utf-8").read(), ""]
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md))

    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            doc.add_heading(b[1], level=1)
        elif kind == "h2":
            doc.add_heading(b[1], level=2)
        elif kind == "p":
            doc.add_paragraph(b[1].replace("**", ""))
        elif kind == "bullets":
            for x in b[1]:
                doc.add_paragraph(x.replace("**", ""), style="List Bullet")
        elif kind == "fig":
            doc.add_picture(b[1], width=Inches(6.0))
            cap = doc.add_paragraph(b[2])
            cap.runs[0].italic = True
        elif kind == "table":
            t = b[1]
            doc.add_paragraph(t["title"]).runs[0].bold = True
            table = doc.add_table(rows=1, cols=len(t["columns"]))
            table.style = "Table Grid"
            for j, (_, head, _) in enumerate(t["columns"]):
                table.rows[0].cells[j].text = head
                table.rows[0].cells[j].paragraphs[0].runs[0].bold = True
            for i in range(len(t["rows"])):
                cells = table.add_row().cells
                for j, (key, _, _) in enumerate(t["columns"]):
                    run = cells[j].paragraphs[0].add_run(t["cell"](i, key))
                    run.bold = (i, key) in t["bold"]
                    run.font.size = Pt(8)
            doc.add_paragraph("")
    doc.save(docx_path)


# ---------------------------------------------------------------- main

def main():
    parser = argparse.ArgumentParser(description="Compare RF / XGBoost / LightGBM / CatBoost / ensemble.")
    parser.add_argument("--quick", action="store_true", help="tiny smoke-test configuration")
    parser.add_argument("--out", default="results", help="results root (default: results)")
    parser.add_argument("--models-dir", default="models", help="where trained models go (default: models)")
    parser.add_argument("--real-root", default="rf_dataset", help="labelled real WAVs (lora_hop/, noise/)")
    parser.add_argument("--no-cache", action="store_true", help="regenerate datasets even if cached")
    args = parser.parse_args()
    cfg = QUICK if args.quick else FULL
    apply_style()
    started = time.time()
    saved = []

    out_models = os.path.join(args.out, "models")
    fig_dir = os.path.join(args.out, "report", "figures")
    tab_dir = os.path.join(args.out, "report", "tables")
    for d in (out_models, fig_dir, tab_dir, args.models_dir):
        os.makedirs(d, exist_ok=True)

    classes = sorted(dsp.CLASSES)          # LabelEncoder order
    encoder = LabelEncoder().fit(dsp.CLASSES)
    deployed = joblib.load("rf_classifier_model.joblib")
    deployed.n_jobs = 1   # in memory only (as --stream does); every model is timed single-threaded
    base = ["rf", "xgb", "lgbm", "cat"]

    # ---- data for every seed (cached in results/models/cache/, keyed by seed, config and DATA_VERSION)
    data = {}
    cache_dir = os.path.join(out_models, "cache")
    os.makedirs(cache_dir, exist_ok=True)
    key = (f"v{DATA_VERSION}_tr{cfg['train_per_class']}_te{cfg['test_per_step']}_e{cfg['e2e_per_step']}_"
           f"snr{min(cfg['snrs'])}to{max(cfg['snrs'])}n{len(cfg['snrs'])}")
    for seed in cfg["seeds"]:
        t0 = time.time()
        path = os.path.join(cache_dir, f"data_seed{seed}_{key}.joblib")
        if os.path.exists(path) and not args.no_cache:
            data[seed] = joblib.load(path)
            source = "cache"
        else:
            Xtr, ytr = make_train(seed, cfg)
            Xte, yte, snr = make_test(seed, cfg)
            e2e = make_e2e(seed, cfg)
            data[seed] = (Xtr, ytr, Xte, yte, snr, e2e)
            joblib.dump(data[seed], path)
            source = "generated"
        print(f"[data] seed {seed}: train {len(data[seed][1])}, test {len(data[seed][3])}, "
              f"end-to-end {len(data[seed][5])} ({source}, {time.time() - t0:.0f} s)", flush=True)

    # ---- tuning on seed-0 training data only
    Xtr0, ytr0 = data[cfg["seeds"][0]][:2]
    y0 = encoder.transform(ytr0)
    folds = StratifiedKFold(cfg["folds"], shuffle=True, random_state=0)
    tuned = {}
    for m in base:
        t0 = time.time()
        search = RandomizedSearchCV(BUILDERS[m](seed=0), PARAM_SPACES[m], n_iter=cfg["trials"], cv=folds,
                                    scoring="f1_macro", n_jobs=-1, random_state=0, refit=False)
        search.fit(Xtr0, y0)
        best = search.best_index_
        params = {k: (v.item() if hasattr(v, "item") else v) for k, v in search.best_params_.items()}
        tuned[m] = {"params": params, "cv_mean": float(search.cv_results_["mean_test_score"][best]),
                    "cv_std": float(search.cv_results_["std_test_score"][best]),
                    "tuning_s": time.time() - t0,
                    "trials": [float(x) for x in search.cv_results_["mean_test_score"]]}
        path = os.path.join(out_models, f"cv_results_{m}.csv")
        with open(path, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["trial", "params", "cv_macro_f1_mean", "cv_macro_f1_std"])
            for i, (p, mu, sd) in enumerate(zip(search.cv_results_["params"],
                                               search.cv_results_["mean_test_score"],
                                               search.cv_results_["std_test_score"])):
                w.writerow([i, json.dumps({k: (v.item() if hasattr(v, "item") else v) for k, v in p.items()}),
                            mu, sd])
        saved.append(path)
        print(f"[tune] {m}: CV macro-F1 {tuned[m]['cv_mean']:.4f} ± {tuned[m]['cv_std']:.4f} "
              f"({tuned[m]['tuning_s']:.0f} s) {params}", flush=True)
    params_all = {m: tuned[m]["params"] for m in base}
    for name, cal in (("ensemble", True), ("ensemble_uncal", False)):
        t0 = time.time()
        scores = cross_val_score(build_ensemble(params_all, 0, cal), Xtr0, y0, cv=folds,
                                 scoring="f1_macro", n_jobs=-1)
        tuned[name] = {"params": {"members": "xgb + lgbm + rf (tuned params)",
                                  "calibration": "isotonic, 5-fold" if cal else "none", "voting": "soft"},
                       "cv_mean": float(scores.mean()), "cv_std": float(scores.std()),
                       "tuning_s": time.time() - t0, "trials": []}
        print(f"[cv]   {name}: CV macro-F1 {scores.mean():.4f} ± {scores.std():.4f}", flush=True)

    def build(name, seed):
        if name in BUILDERS:
            return BUILDERS[name](params_all[name], seed)
        return build_ensemble(params_all, seed, calibrated=(name == "ensemble"))

    # ---- fit + evaluate every model on every seed
    names = base + ["ensemble", "ensemble_uncal"]
    res = {m: [] for m in names + ["rf_deployed"]}
    keep = {}
    for seed in cfg["seeds"]:
        Xtr, ytr, Xte, yte, snr, e2e = data[seed]
        for m in names:
            est = build(m, seed)
            t0 = time.time()
            est.fit(Xtr, encoder.transform(ytr))
            fit_s = time.time() - t0
            model = LabeledModel(est, encoder)
            r = evaluate(model, Xte, yte, snr, classes, cfg)
            r.update(evaluate_e2e(model, e2e, cfg))
            r["train_s"] = fit_s
            res[m].append(r)
            if seed == cfg["seeds"][0]:
                keep[m] = model
            print(f"[eval] seed {seed} {m:15s} macro-F1 {r['macro_f1']:.3f}  0-10 dB {r['f1_0_10']:.3f}  "
                  f"e2e 0-10 dB {r['e2e_0_10']:.3f}  fit {fit_s:.1f} s", flush=True)
        r = evaluate(deployed, Xte, yte, snr, classes, cfg)
        r.update(evaluate_e2e(deployed, e2e, cfg))
        r["train_s"] = float("nan")
        res["rf_deployed"].append(r)
    keep["rf_deployed"] = deployed

    # ---- save models (seed-0 fits), sizes, latency, permutation importance
    Xte0, yte0, snr0 = data[cfg["seeds"][0]][2:5]
    hi = snr0 >= 0
    extra = {}
    for m, model in keep.items():
        if m == "rf_deployed":
            size = os.path.getsize("rf_classifier_model.joblib")
        else:
            meta = {"name": m, "classes": list(encoder.classes_), "features": dsp.FEATURE_NAMES,
                    "params": tuned[m]["params"], "cv_macro_f1_mean": tuned[m]["cv_mean"],
                    "cv_macro_f1_std": tuned[m]["cv_std"], "seed": cfg["seeds"][0],
                    "snr_definition": dsp.SNR_DEFINITION, "created": datetime.now().isoformat(),
                    "quick": args.quick}
            paths = save_model(SAVED_NAME[m], model, meta, args.models_dir)
            saved += list(paths)
            size = os.path.getsize(paths[0])
        extra[m] = {"size_mb": size / 1e6, **latency_ms(model, Xte0),
                    "importance": permutation_drop(model, Xte0[hi], yte0[hi])}

    real = real_eval(keep, args.real_root) if os.path.isdir(args.real_root) else None

    # ---- aggregate
    order = [m for m in MODEL_ORDER if m in res]

    def ms(m, key):
        vals = np.array([r[key] for r in res[m]], dtype=float)
        return (float(np.nanmean(vals)), float(np.nanstd(vals))) if not np.all(np.isnan(vals)) else (float("nan"), 0.0)

    summary = {m: {k: ms(m, k) for k in ("macro_f1", "balanced_accuracy", "pd_at_pfa1", "log_loss", "brier",
                                          "ece", "f1_0_10", "snr_f1_90", "e2e_0_10", "noise_false_alarm",
                                          "train_s")} for m in order}

    # CSVs in results/models
    path = os.path.join(out_models, "f1_vs_snr.csv")
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["model", "seed", "evaluation"] + [f"{s} dB" for s in cfg["snrs"]])
        for m in order:
            for seed, r in zip(cfg["seeds"], res[m]):
                w.writerow([m, seed, "classifier_macro_f1"] + list(r["f1_vs_snr"]))
                w.writerow([m, seed, "end_to_end_pd_x_correct"] + list(r["e2e_vs_snr"]))
                w.writerow([m, seed, "end_to_end_pd"] + list(r["pd_vs_snr"]))
    saved.append(path)
    path = os.path.join(out_models, "comparison.csv")
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        keys = list(summary[order[0]].keys())
        w.writerow(["model"] + [f"{k}_{s}" for k in keys for s in ("mean", "std")]
                   + ["latency_b1_ms", "latency_b10_ms", "size_mb"])
        for m in order:
            w.writerow([m] + [v for k in keys for v in summary[m][k]]
                       + [extra[m]["latency_b1"], extra[m]["latency_b10"], extra[m]["size_mb"]])
    saved.append(path)

    # ---- figures
    snrs = cfg["snrs"]
    figs = {}
    figs["f1"] = fig_curve(snrs, {m: [r["f1_vs_snr"] for r in res[m]] for m in order},
                           "Macro-F1", "Classifier only (true signal slice)",
                           os.path.join(fig_dir, "fig01_f1_vs_snr"), hline=0.9)
    figs["e2e"] = fig_curve(snrs, {m: [r["e2e_vs_snr"] for r in res[m]] for m in order},
                            "Pd × correct label", "End-to-end through detect_regions",
                            os.path.join(fig_dir, "fig02_end_to_end_vs_snr"), hline=0.9)
    cms = {}
    for m in order:
        r0 = res[m][0]
        cm = np.zeros((len(CLASS_ORDER), len(CLASS_ORDER)))
        for t, p in zip(yte0, r0["pred"]):
            cm[CLASS_ORDER.index(t), CLASS_ORDER.index(p)] += 1
        cms[m] = cm / cm.sum(axis=1, keepdims=True)
    figs["cm"] = fig_confusions(cms, CLASS_ORDER, os.path.join(fig_dir, "fig03_confusion_matrices"))
    rocs = {}
    for m in order:
        s = res[m][0]["proba"][:, classes.index(POSITIVE)]
        fpr, tpr, _ = roc_curve(yte0 == POSITIVE, s)
        rocs[m] = (fpr, tpr)
    figs["roc"] = fig_roc(rocs, os.path.join(fig_dir, "fig04_roc_lora"))
    rel = {}
    for m in order:
        proba = res[m][0]["proba"]
        conf = proba.max(axis=1)
        correct = np.array(classes)[proba.argmax(axis=1)] == yte0
        edges = np.linspace(0, 1, 11)
        cx, ay = [], []
        for lo, hi_ in zip(edges[:-1], edges[1:]):
            sel = (conf > lo) & (conf <= hi_)
            if sel.sum() >= 5:
                cx.append(conf[sel].mean())
                ay.append(correct[sel].mean())
        rel[m] = (cx, ay, res[m][0]["ece"])
    figs["rel"] = fig_reliability(rel, os.path.join(fig_dir, "fig05_reliability"))
    figs["lat"] = fig_latency({m: (extra[m]["latency_b1"], summary[m]["f1_0_10"][0]) for m in order},
                              os.path.join(fig_dir, "fig06_latency_vs_f1"))
    f1mat = np.array([[np.mean([r["f1"][c] for r in res[m]]) for c in CLASS_ORDER] for m in order])
    figs["heat"] = fig_heatmap(f1mat, [MODEL_LABELS[m] for m in order], [SHORT_CLASS[c] for c in CLASS_ORDER],
                               "F1 (mean of seeds)", os.path.join(fig_dir, "fig07_per_class_f1"),
                               vmin=0, vmax=1)
    impmat = np.array([extra[m]["importance"] for m in order]).T
    figs["imp"] = fig_heatmap(impmat, dsp.FEATURE_NAMES, [MODEL_LABELS[m] for m in order],
                              "Macro-F1 drop when permuted", os.path.join(fig_dir, "fig08_feature_importance"),
                              fmt="{:.3f}", cmap="magma", height=5.2)
    if real:
        rn = [m for m in order]
        synth = []
        for m in rn:
            p0 = res[m][0]["pred"]
            synth.append(balanced_accuracy_score(yte0 == POSITIVE, p0 == POSITIVE))
        figs["real"] = fig_bars(rn, synth, [real["results"][m]["balanced_accuracy"] for m in rn],
                                os.path.join(fig_dir, "fig09_synthetic_vs_real"))
    for p in figs.values():
        saved += p
    # raw plots: per-seed curves and tuning traces
    fig = Figure(figsize=(DOUBLE_COL, 3.4))
    ax = fig.add_subplot()
    for m in order:
        for k, r in enumerate(res[m]):
            ax.plot(snrs, r["f1_vs_snr"], color=MODEL_COLORS[m], lw=0.8,
                    label=MODEL_LABELS[m] if k == 0 else None)
    ax.set_xlabel("In-band SNR (dB)")
    ax.set_ylabel("Macro-F1")
    ax.set_title("Classifier-only macro-F1, every seed")
    legend_outside(ax, fig)
    saved += save_figure(fig, os.path.join(out_models, "raw_f1_vs_snr_per_seed"))
    fig = Figure(figsize=(DOUBLE_COL, 3.0))
    ax = fig.add_subplot()
    for m in base:
        ax.plot(range(1, len(tuned[m]["trials"]) + 1), sorted(tuned[m]["trials"]), color=MODEL_COLORS[m],
                label=MODEL_LABELS[m], marker="o", markersize=3)
    ax.set_xlabel("Trial (sorted by CV score)")
    ax.set_ylabel("CV macro-F1")
    ax.set_title("Random-search trials (seed-0 training set)")
    legend_outside(ax, fig, ncol=4)
    saved += save_figure(fig, os.path.join(out_models, "raw_tuning_trials"))

    # ---- tables
    tables = {}
    rows = [{"model": MODEL_LABELS[m], "f1": summary[m]["macro_f1"], "bacc": summary[m]["balanced_accuracy"],
             "pd": summary[m]["pd_at_pfa1"], "snr90": summary[m]["snr_f1_90"], "f1lo": summary[m]["f1_0_10"],
             "e2e": summary[m]["e2e_0_10"], "lat1": extra[m]["latency_b1"], "lat10": extra[m]["latency_b10"],
             "size": extra[m]["size_mb"], "train": summary[m]["train_s"] if m != "rf_deployed" else "n/a"}
            for m in order]
    tables[1] = write_table("table1_overall", "Table 1. Overall comparison (synthetic test set, -10 to +30 dB; "
                            "mean ± std over seeds)",
                            [("model", "Model", None), ("f1", "Macro-F1", "max"), ("bacc", "Balanced acc.", "max"),
                             ("pd", "Pd @ Pfa=1% (LORA)", "max"), ("snr90", "SNR for 90% F1 (dB)", "min"),
                             ("f1lo", "Macro-F1 0–10 dB", "max"), ("e2e", "End-to-end 0–10 dB", "max"),
                             ("lat1", "Latency b1 (ms/det)", "min"), ("lat10", "Latency b10 (ms/det)", "min"),
                             ("size", "Size (MB)", "min"), ("train", "Train time (s)", "min")], rows, tab_dir)
    rows = []
    for c in CLASS_ORDER:
        for m in order:
            rows.append({"cls": c, "model": MODEL_LABELS[m],
                         "p": ms_class(res[m], "precision", c), "r": ms_class(res[m], "recall", c),
                         "f": ms_class(res[m], "f1", c)})
    tables[2] = write_table("table2_per_class", "Table 2. Per-class precision / recall / F1 (best per class in bold)",
                            [("cls", "Class", None), ("model", "Model", None), ("p", "Precision", "max"),
                             ("r", "Recall", "max"), ("f", "F1", "max")], rows, tab_dir, group_key="cls")
    rows = []
    for m in order:
        row = {"model": MODEL_LABELS[m]}
        for s in TABLE_SNRS:
            if s in snrs:
                k = snrs.index(s)
                vals = [r["f1_vs_snr"][k] for r in res[m]]
                row[f"s{s}"] = (float(np.mean(vals)), float(np.std(vals)))
            else:
                row[f"s{s}"] = "n/a"
        rows.append(row)
    tables[3] = write_table("table3_f1_at_snr", "Table 3. Classifier-only macro-F1 at selected in-band SNRs",
                            [("model", "Model", None)] + [(f"s{s}", f"{s} dB", "max") for s in TABLE_SNRS],
                            rows, tab_dir)
    rows = [{"model": MODEL_LABELS[m],
             "params": ", ".join(f"{k}={sig3(v) if isinstance(v, float) else v}" for k, v in tuned[m]["params"].items()),
             "cv": (tuned[m]["cv_mean"], tuned[m]["cv_std"]), "tune": tuned[m]["tuning_s"]}
            for m in names]
    tables[4] = write_table("table4_hyperparameters", f"Table 4. Tuned hyperparameters ({cfg['trials']}-trial random "
                            f"search, {cfg['folds']}-fold stratified CV on the seed-0 training set)",
                            [("model", "Model", None), ("params", "Best hyperparameters", None),
                             ("cv", "CV macro-F1", "max"), ("tune", "Tuning time (s)", "min")], rows, tab_dir)
    rows = [{"model": MODEL_LABELS[m], "ll": summary[m]["log_loss"], "brier": summary[m]["brier"],
             "ece": summary[m]["ece"], "fa": summary[m]["noise_false_alarm"]} for m in order]
    tables[5] = write_table("table5_calibration", "Table 5. Probability quality and end-to-end false alarms",
                            [("model", "Model", None), ("ll", "Log loss", "min"), ("brier", "Brier score", "min"),
                             ("ece", "ECE", "min"), ("fa", "Noise captures with a non-NOISE label", "min")],
                            rows, tab_dir)
    for t in tables.values():
        saved += [t["csv"], t["md"], t["tex"]]

    # ---- recommendation
    within = [m for m in order if m != "rf_deployed" and extra[m]["latency_b1"] < BLOCK_BUDGET_MS]
    best_f1 = max(summary[m]["f1_0_10"][0] for m in within)
    close = [m for m in within if summary[m]["f1_0_10"][0] >= best_f1 - 0.01]
    rec = min(close, key=lambda m: extra[m]["latency_b1"])
    rec_text = (f"{MODEL_LABELS[rec]} (--model {SAVED_NAME[rec] if rec != 'rf' else 'rf_tuned'}): macro-F1 "
                f"{summary[rec]['f1_0_10'][0]:.3f} at 0–10 dB SNR (best {best_f1:.3f}; models within 0.01 of the "
                f"best: {', '.join(MODEL_LABELS[m] for m in close)}), {extra[rec]['latency_b1']:.2f} ms per "
                f"detection, far inside the {BLOCK_BUDGET_MS} ms block budget. Chosen as the fastest of the "
                f"models within 0.01 macro-F1 of the best.")

    # ---- report
    report_dir = os.path.join(args.out, "report")
    blocks = build_report(cfg, args, order, names, summary, extra, tuned, res, tables, figs, real, rec,
                          rec_text, fig_dir, time.time() - started, cms)
    md_path = os.path.join(report_dir, "model_comparison_report.md")
    docx_path = os.path.join(report_dir, "model_comparison_report.docx")
    render_report(blocks, md_path, docx_path)
    saved += [md_path, docx_path]

    print("\nRECOMMENDATION: " + rec_text)
    print(f"\nSaved files ({len(saved)}):")
    for p in saved:
        print("  " + os.path.normpath(p))
    print(f"\nTotal time {time.time() - started:.0f} s")


def ms_class(runs, key, cls):
    vals = [r[key][cls] for r in runs]
    return (float(np.mean(vals)), float(np.std(vals)))


def build_report(cfg, args, order, names, summary, extra, tuned, res, tables, figs, real, rec, rec_text,
                 fig_dir, elapsed, cms):
    L = MODEL_LABELS
    n_seeds = len(cfg["seeds"])
    env = environment()
    f1 = {m: summary[m]["macro_f1"][0] for m in order}
    lo = {m: summary[m]["f1_0_10"][0] for m in order}
    e2e = {m: summary[m]["e2e_0_10"][0] for m in order}
    snr90 = {m: summary[m]["snr_f1_90"][0] for m in order}
    best_all = max(names, key=f1.get)
    best_lo = max(names, key=lo.get)
    fastest = min(names, key=lambda m: extra[m]["latency_b1"])
    slowest = max(names, key=lambda m: extra[m]["latency_b1"])
    ece_cal, ece_unc = summary["ensemble"]["ece"][0], summary["ensemble_uncal"]["ece"][0]
    ll_cal, ll_unc = summary["ensemble"]["log_loss"][0], summary["ensemble_uncal"]["log_loss"][0]
    k0 = cfg["snrs"].index(0) if 0 in cfg["snrs"] else None
    gap = ""
    if k0 is not None:
        c0 = np.mean([r["f1_vs_snr"][k0] for r in res[best_lo]])
        e0 = np.mean([r["e2e_vs_snr"][k0] for r in res[best_lo]])
        pd0 = np.mean([r["pd_vs_snr"][k0] for r in res[best_lo]])
        gap = (f"At 0 dB in-band SNR, {L[best_lo]} reaches macro-F1 {c0:.3f} when given the true signal slice, "
               f"but only {e0:.3f} end-to-end (detected and correctly labelled); the detector finds {pd0:.0%} of "
               f"signals at that SNR, so detection, not classification, limits low-SNR performance.")
    spread = max(lo[m] for m in names) - min(lo[m] for m in names)
    findings = [
        f"Overall macro-F1 (all SNRs, -10 to +30 dB): best {L[best_all]} {f1[best_all]:.3f}; range across tuned "
        f"models {min(f1[m] for m in names):.3f}–{max(f1[m] for m in names):.3f}. The deployed RF scores "
        f"{f1['rf_deployed']:.3f} on the same test set.",
        f"At 0–10 dB SNR the best model is {L[best_lo]} (macro-F1 {lo[best_lo]:.3f}); the spread across tuned "
        f"models there is {spread:.3f}, " + ("so the choice of model matters little at these SNRs."
                                               if spread < 0.02 else "a meaningful difference."),
        "SNR needed for 90 % macro-F1: " + ", ".join(f"{L[m]} {snr90[m]:.1f} dB" if not math.isnan(snr90[m])
                                                     else f"{L[m]} not reached" for m in order) + ".",
        f"Calibration: the calibrated ensemble has ECE {ece_cal:.3f} and log loss {ll_cal:.3f} versus "
        f"{ece_unc:.3f} and {ll_unc:.3f} uncalibrated — "
        + ("calibration improves both." if ece_cal < ece_unc and ll_cal < ll_unc else
           "calibration lowers ECE but raises log loss (isotonic calibration assigns near-zero probability "
           "to some true classes, which log loss punishes heavily)." if ece_cal < ece_unc else
           "calibration lowers log loss but not ECE." if ll_cal < ll_unc else
           "calibration improves neither here."),
        f"Latency per detection (batch of 1): fastest {L[fastest]} {extra[fastest]['latency_b1']:.2f} ms, slowest "
        f"{L[slowest]} {extra[slowest]['latency_b1']:.2f} ms; all are far below the {BLOCK_BUDGET_MS} ms "
        f"block budget, where feature extraction and detection dominate.",
    ]
    if gap:
        findings.append(gap)
    hi_cls = {c: np.mean([r["e2e_high_per_class"][c] for r in res[best_lo]]) for c in CLASS_ORDER if c != "NOISE"}
    k20 = [k for k, s in enumerate(cfg["snrs"]) if s >= 20]
    clf_hi = np.mean([np.mean([r["f1_vs_snr"][k] for k in k20]) for r in res[best_lo]])
    e2e_hi = np.mean([np.mean([r["e2e_vs_snr"][k] for k in k20]) for r in res[best_lo]])
    findings.append(
        f"Even at >= 20 dB, where every signal is detected, {L[best_lo]} labels {e2e_hi:.0%} of detected signals "
        f"correctly end-to-end versus classifier-only macro-F1 {clf_hi:.3f}. Per class (end-to-end, >= 20 dB): "
        + ", ".join(f"{c} {v:.0%}" for c, v in hi_cls.items())
        + ". The classifiers learn from the true signal slice, but the pipeline hands them its own bandwidth "
          "estimate and burst window; where those differ (e.g. a narrow Welch bandwidth on a line-spectrum FM "
          "signal, or a single telemetry burst), the features no longer match training. Training on "
          "pipeline-extracted features would close this gap; the choice of classifier does not.")
    if real:
        findings.append("Real recordings: " + ", ".join(
            f"{L[m]} {real['results'][m]['balanced_accuracy']:.2f}" for m in order) + " balanced accuracy (LORA vs rest).")

    counts = ", ".join(f"{c}: {cfg['train_per_class']}" for c in CLASS_ORDER)
    n_test = cfg["test_per_step"] * len(cfg["snrs"])
    blocks = [
        ("h1", "RF signal classifier comparison"),
        ("p", f"Generated {datetime.now():%Y-%m-%d %H:%M} by compare_models.py"
              + (" (QUICK smoke-test configuration — not for reporting)." if args.quick else ".")),
        ("h2", "Objective"),
        ("p", "Compare Random Forest, XGBoost, LightGBM, CatBoost and a soft-voting ensemble for classifying "
              "detected RF signals in the RTL-SDR drone threat pipeline, on identical features, data and tuning "
              "budgets, and recommend a model for live use within the 500 ms per-block budget."),
        ("h2", "Experimental setup"),
        ("bullets", [
            f"Classes: {', '.join(CLASS_ORDER)}.",
            f"Training set per seed: {counts} snippets (15–500 ms), in-band SNR uniform 0–30 dB.",
            f"Test set per seed: {cfg['test_per_step']} snippets per class per SNR step, SNR steps "
            f"{', '.join(str(s) for s in cfg['snrs'])} dB ({n_test} per class). End-to-end set: "
            f"{cfg['e2e_per_step']} captures (0.5 s) per class per SNR step.",
            "SNR definition: " + dsp.SNR_DEFINITION,
            "Randomised per sample: frequency offset (±300 kHz), carrier phase, amplitude (±20 dB), duration, "
            "and for LORA the spreading factor (7–9) and bandwidth (125/250/500 kHz).",
            "Features (11, the pipeline's own code in dsp.py): " + ", ".join(dsp.FEATURE_NAMES) + ".",
            f"Split: train, test and end-to-end sets are generated independently (separate RNG streams); "
            f"seeds {cfg['seeds']} each draw their own sets and model seed.",
            f"Tuning: {cfg['trials']}-trial random search per model, {cfg['folds']}-fold stratified CV (macro-F1) on "
            f"the seed-0 training set only; best parameters reused for all seeds. Ensemble members use their tuned "
            f"parameters and are calibrated (CalibratedClassifierCV, isotonic, 5-fold) before soft voting; an "
            f"uncalibrated ensemble is reported for comparison.",
            "Evaluations: classifier-only (features from the true signal slice, as in training) and end-to-end "
            "(a 0.5 s capture the signal spans — continuous classes on throughout, bursty classes with their own "
            "gaps — passed through detect_regions and the pipeline's own feature extraction; success = detected "
            "and correctly labelled; NOISE success = no non-NOISE label).",
            "Pd at Pfa = 1 %: LORA vs rest using P(LORA), threshold at the 99th percentile of non-LORA scores.",
            "Latency: median predict_proba time per detection (batch 1 and 10), excluding feature extraction.",
            "Hardware/software: " + "; ".join(f"{k} {v}" for k, v in env.items()) + ".",
            f"Total run time: {elapsed / 60:.1f} min.",
        ]),
        ("h2", "Results"),
    ]
    captions = [
        ("f1", "Figure 1. Classifier-only macro-F1 versus in-band SNR (line: mean of 3 seeds; band: min–max). "
               "Dotted line: 90 %."),
        ("e2e", "Figure 2. End-to-end success (detected × correctly labelled) versus in-band SNR, averaged over "
                "the signal classes."),
        ("cm", "Figure 3. Normalised confusion matrices on the seed-0 test set (all SNR steps)."),
        ("roc", "Figure 4. Pd versus Pfa for LORA vs rest (seed 0); dotted line at Pfa = 1 %."),
        ("rel", "Figure 5. Reliability diagrams (top-label confidence vs accuracy, seed 0)."),
        ("lat", "Figure 6. Inference latency (batch of 1) versus macro-F1 at 0–10 dB SNR."),
        ("heat", "Figure 7. Per-class F1 (mean of seeds)."),
        ("imp", "Figure 8. Permutation importance (macro-F1 drop, seed-0 test set, SNR ≥ 0 dB)."),
        ("real", "Figure 9. LORA-vs-rest balanced accuracy on the synthetic test set and on real recordings."),
    ]
    snrs = cfg["snrs"]
    pd_mean = np.mean([r["pd_vs_snr"] for r in res[best_lo]], axis=0)
    pd50 = next((s for s, p in zip(snrs, pd_mean) if p >= 0.5), None)
    pd90 = next((s for s, p in zip(snrs, pd_mean) if p >= 0.9), None)
    conf_pairs = {}
    for m in names:
        cm = cms[m]
        for i in range(len(CLASS_ORDER)):
            for j in range(len(CLASS_ORDER)):
                if i != j:
                    conf_pairs[(i, j)] = conf_pairs.get((i, j), 0) + cm[i, j] / len(names)
    (ci, cj), cval = max(conf_pairs.items(), key=lambda kv: kv[1])
    worst = {m: min(CLASS_ORDER, key=lambda c: np.mean([r["f1"][c] for r in res[m]])) for m in names}
    worst_counts = {}
    for c in worst.values():
        worst_counts[c] = worst_counts.get(c, 0) + 1
    mean_imp = {f: np.mean([extra[m]["importance"][k] for m in names]) for k, f in enumerate(dsp.FEATURE_NAMES)}
    top = sorted(mean_imp, key=mean_imp.get, reverse=True)[:3]
    notes = {
        "f1": f"{L[best_lo]} leads at 0–10 dB; the SNR for 90 % macro-F1 ranges from "
              f"{np.nanmin(list(snr90.values())):.1f} to {np.nanmax(list(snr90.values())):.1f} dB across models.",
        "e2e": f"For {L[best_lo]} the detector finds half the signals from about {pd50} dB and 90 % from about "
               f"{pd90} dB in-band SNR; below that the end-to-end curve is limited by detection, not by the "
               f"classifier.",
        "cm": f"The largest confusion, averaged over the tuned models, is true {CLASS_ORDER[ci]} predicted as "
              f"{CLASS_ORDER[cj]} ({cval:.0%} of that class), pooled over all SNR steps.",
        "roc": "Pd at Pfa = 1 %: " + ", ".join(f"{L[m]} {summary[m]['pd_at_pfa1'][0]:.3f}" for m in order) + ".",
        "rel": f"ECE ranges from {min(summary[m]['ece'][0] for m in order):.3f} to "
               f"{max(summary[m]['ece'][0] for m in order):.3f}.",
        "lat": f"Classifier latency ranges from {min(extra[m]['latency_b1'] for m in order):.2f} to "
               f"{max(extra[m]['latency_b1'] for m in order):.1f} ms per detection (batch of 1), all inside the "
               f"{BLOCK_BUDGET_MS} ms block budget.",
        "heat": "Weakest class per tuned model: " + ", ".join(f"{c} ({n} of {len(names)} models)"
                                                         for c, n in sorted(worst_counts.items(),
                                                                            key=lambda kv: -kv[1])) + ".",
        "imp": "Most important features on average (macro-F1 drop when permuted): "
               + ", ".join(f"{f} ({mean_imp[f]:.3f})" for f in top) + ".",
        "real": "Real-data accuracy should be read with the clip count in mind (see Limitations).",
    }
    for key, cap in captions:
        if key in figs:
            blocks += [("fig", figs[key][0], cap), ("p", notes[key])]
    scenario = sorted(glob.glob(os.path.join(fig_dir, "fig1[0-9]_scenario_*.png")))
    for i, path in enumerate(scenario):
        name = os.path.basename(path)[len("fig10_scenario_"):-4].replace("_", " ")
        blocks += [("fig", path, f"Figure {10 + i}. Scenario '{name}': spectrum (Welch average, max-hold, "
                                 f"thresholds, detections) and waterfall with detected bursts boxed."),
                   ("p", "Produced by spectrum_report.py with the deployed pipeline.")]
    for k in sorted(tables):
        blocks.append(("table", tables[k]))
    blocks += [
        ("h2", "Key findings"),
        ("bullets", findings),
        ("h2", "Limitations"),
        ("bullets", [
            "All training and test data are synthetic. Real emitters (multipath, oscillator drift, front-end "
            "non-linearity, other users of 865–867 MHz) are not represented, so absolute scores overstate real "
            "performance (domain gap).",
            ("No labelled real recordings were found in rf_dataset/, so there is no real-data evaluation."
             if not real else f"Real-data evaluation uses only {len(real['clips'])} clips; treat it as indicative."),
            "End-to-end captures contain one signal each; crowded bands with overlapping emitters are not tested.",
            "Latency is measured on this machine for the classifier only; detection and feature extraction "
            "(tens to hundreds of ms per block) dominate the pipeline's time.",
            "The deployed RF was trained by train_rf_classifier.py on a different synthetic generator (whole-band "
            "SNR for most classes), so it is a reference, not a tuned competitor.",
        ]),
        ("h2", "Recommended model"),
        ("p", rec_text),
        ("p", "The default pipeline stays on the deployed RF (--model rf); other models load with "
              "--model rf_tuned/xgb/lgbm/cat/ensemble."),
    ]
    return blocks


if __name__ == "__main__":
    sys.exit(main())
